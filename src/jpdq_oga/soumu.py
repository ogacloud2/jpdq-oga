"""Importer for the 総務省 全国地方公共団体コード workbooks.

Two files are published at https://www.soumu.go.jp/denshijiti/code.html:

``都道府県コード及び市区町村コード`` (current codes)
    Sheets ``R6.1.1現在の団体`` and ``R6.1.1政令指定都市`` (the second lists the wards of designated
    cities).  Header row ``団体コード | 都道府県名(漢字) | 市区町村名(漢字) | ...``; the prefecture's own
    row has no city name.  The *union* of all sheets is the set of codes in force on the sheet date.

``「都道府県コード及び市区町村コード」の改正一覧表（平成17年4月1日以降）`` (revisions)
    Sheet ``改正一覧表``, three header rows, then blocks of rows.  A block starts at a row whose
    ``都道府県名`` (column E) is filled.  Per row: before-code (F) and name (G), the change category (I:
    ``欠番`` / ``削除`` retire the before-code, ``新設`` creates the after-code (J), ``〃`` repeats the row
    above, empty = the code is unchanged, e.g. a rename), and the effective date (N).  The date is
    written once per block, on whichever row carries it, as ``R6.1.1`` / ``H30.10.1``, a date cell, or
    an Excel serial number.  Codes may be stored as integers (leading zero lost).

The structure above was checked against the files published on 2026-10-15; the importer refuses
workbooks that do not look like that instead of guessing.  Not covered: renames over time (the current
name is kept), and revisions before 平成17年4月1日 (not in the file).
"""

from __future__ import annotations

import datetime as dt
import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from .reference import ReferenceDataError
from .wareki import parse_wareki

CURRENT_HEADER = "団体コード"
REVISION_SHEET = "改正一覧表"
RETIRING = ("欠番", "削除")
NEW = "新設"
DITTO = "〃"

# 0-based column indexes of 改正一覧表
COL_PREFECTURE, COL_BEFORE_CODE, COL_BEFORE_NAME, COL_CATEGORY, COL_AFTER_CODE, COL_AFTER_NAME, COL_DATE = (
    4, 5, 6, 8, 9, 10, 13,
)  # fmt: skip


@dataclass(frozen=True)
class CurrentCode:
    code: str
    prefecture: str
    city: str | None


@dataclass(frozen=True)
class RevisionEvent:
    kind: str  # "retire" | "new"
    code: str
    prefecture: str | None
    city: str | None
    effective: dt.date


@dataclass
class Revisions:
    events: list[RevisionEvent]
    skipped: list[str] = field(default_factory=list)  # human-readable reasons for rows that were not used


def _openpyxl() -> Any:
    try:
        import openpyxl
    except ImportError as exc:  # pragma: no cover - depends on the optional extra
        raise ReferenceDataError(
            "reading the 総務省 Excel files needs openpyxl: pip install 'jpdq-oga[excel]'"
        ) from exc
    return openpyxl


def code6(value: Any) -> str | None:
    """A 6-digit code from a string or an integer cell whose leading zero was lost; else ``None``."""
    if value is None or isinstance(value, bool):
        return None
    text = str(int(value)) if isinstance(value, float) and value == int(value) else str(value).strip()
    return text.zfill(6) if re.fullmatch(r"[0-9]{1,6}", text) else None


def parse_effective_date(value: Any) -> dt.date | None:
    """``R6.1.1`` / ``H30.10.1`` strings, date cells and Excel serial numbers (1900 date system)."""
    if isinstance(value, dt.datetime):
        return value.date()
    if isinstance(value, dt.date):
        return value
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return dt.date(1899, 12, 30) + dt.timedelta(days=int(value))
    if isinstance(value, str):
        return parse_wareki(value.strip())
    return None


def _clean(value: Any) -> str | None:
    text = None if value is None else str(value).strip()
    return text or None


def read_current_codes(path: str | Path) -> tuple[list[CurrentCode], dt.date | None]:
    """Codes in force (union of all sheets) and the date written in the sheet name (``R6.1.1現在``)."""
    workbook = _openpyxl().load_workbook(path, data_only=True)
    codes: dict[str, CurrentCode] = {}
    as_of: dt.date | None = None
    for sheet in workbook.worksheets:
        rows = sheet.iter_rows(values_only=True)
        header = next(rows, None)
        if not header or _clean(header[0]) != CURRENT_HEADER:
            continue
        match = re.match(r"([RHS][0-9]+\.[0-9]+\.[0-9]+)", sheet.title)
        if match and as_of is None:
            as_of = parse_wareki(match.group(1))
        for row in rows:
            code = code6(row[0])
            if code is None:
                continue
            codes.setdefault(code, CurrentCode(code, _clean(row[1]) or "", _clean(row[2])))
    if not codes:
        raise ReferenceDataError(
            f"{path}: no sheet with the header {CURRENT_HEADER!r} in A1 and 6-digit codes; "
            "is this the 総務省「都道府県コード及び市区町村コード」 workbook?"
        )
    return list(codes.values()), as_of


def read_revisions(path: str | Path) -> Revisions:
    """Retire/new events of the 改正一覧表, each with its effective date."""
    workbook = _openpyxl().load_workbook(path, data_only=True)
    sheet = workbook[REVISION_SHEET] if REVISION_SHEET in workbook.sheetnames else workbook.worksheets[0]
    rows = list(sheet.iter_rows(values_only=True))
    header_at = next(
        (i for i, r in enumerate(rows[:10]) if len(r) > COL_DATE and _clean(r[COL_CATEGORY]) == "コードの改定区分"),
        None,
    )
    if header_at is None or _clean(rows[header_at][COL_DATE]) != "施行年月日":
        raise ReferenceDataError(
            f"{path}: expected the 改正一覧表 layout (header row with 'コードの改定区分' in column I and "
            "'施行年月日' in column N)"
        )
    data = rows[header_at + 3 :]  # two more header rows follow (コード/市区町村名, 名称/ふりがな)

    blocks: list[list[tuple[Any, ...]]] = []
    for row in data:
        row = tuple(row) + (None,) * (COL_DATE + 1 - len(row))
        if all(c is None for c in row):
            continue
        if _clean(row[COL_PREFECTURE]) and (not blocks or blocks[-1]):
            blocks.append([])
        if not blocks:
            raise ReferenceDataError(f"{path}: data does not start with a prefecture name in column E")
        blocks[-1].append(row)

    result = Revisions(events=[])
    for block in blocks:
        prefecture = _clean(block[0][COL_PREFECTURE])
        dates = [d for r in block if r[COL_DATE] not in (None, DITTO) and (d := parse_effective_date(r[COL_DATE]))]
        category: str | None = None
        for row in block:
            cat = _clean(row[COL_CATEGORY])
            if cat == DITTO:
                cat = category
            elif cat is not None:
                category = cat
            if cat not in (*RETIRING, NEW):
                continue
            kind, code, name = (
                ("retire", code6(row[COL_BEFORE_CODE]), _clean(row[COL_BEFORE_NAME]))
                if cat in RETIRING
                else ("new", code6(row[COL_AFTER_CODE]), _clean(row[COL_AFTER_NAME]))
            )
            if code is None:
                continue  # e.g. the "削除" marker row of a block that only lists the after-codes
            if not dates:
                result.skipped.append(f"{kind} {code} ({prefecture} {name}): the block has no effective date")
                continue
            result.events.append(RevisionEvent(kind, code, prefecture, name, dates[0]))
    return result


HistoryRow = tuple[str, str | None, str | None, str | None, str | None]  # code, prefecture, city, from, to


def build_history(current: list[CurrentCode], events: list[RevisionEvent]) -> list[HistoryRow]:
    """``local_gov_codes`` rows: current codes plus the codes retired since 平成17年4月1日.

    ``valid_from`` is the date of the ``新設`` event if the code was created inside the revision window,
    else ``None`` (in force since before the window); ``valid_to`` is the retirement date.
    """
    created = {e.code: e.effective for e in events if e.kind == "new"}
    retired = {e.code: e for e in events if e.kind == "retire"}
    in_force = {c.code for c in current}
    clash = sorted(retired.keys() & in_force)
    if clash:
        raise ReferenceDataError(
            f"codes are both retired by the revisions and listed as current: {clash[:5]} "
            "(are the two files from the same release?)"
        )

    def iso(d: dt.date | None) -> str | None:
        return d.isoformat() if d else None

    rows: list[HistoryRow] = [(c.code, c.prefecture or None, c.city, iso(created.get(c.code)), None) for c in current]
    rows += [(code, e.prefecture, e.city, iso(created.get(code)), iso(e.effective)) for code, e in retired.items()]
    return sorted(rows)
